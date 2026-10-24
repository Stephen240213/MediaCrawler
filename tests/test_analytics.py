"""Analytics contract tests. All database mutations use pytest temporary paths."""
import csv
import io
import json
import sqlite3
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import text

from api.main import app
from api.services import analytics as service
from database import analytics as schema
from database import db_session


@pytest.fixture
def database(tmp_path, monkeypatch):
    path = tmp_path / 'collection.db'
    monkeypatch.setitem(service.sqlite_db_config, 'db_path', str(path))
    yield path


@pytest.fixture
def client(database):
    with TestClient(app) as client:
        yield client


def note(id='a', **kwargs):
    return dict(platform='xhs', content_id=id, author='匿名作者', keyword='旅行',
                published_at='2026-01-01T12:00:00+08:00', collected_at='2026-02-01T12:00:00+08:00', likes='1.2万', **kwargs)


def insert(path, records, kind='content'):
    schema.migrate(path)
    with sqlite3.connect(path) as db:
        for item in records:
            schema.save(db, schema.normalize(item, kind=kind), kind)


@pytest.mark.parametrize('value,expected', [('1.2万',12000),('2.5k',2500),('1,234',1234),('0',0),('',None),(None,None),('None',None),('未知',None),('-1',None),('1亿',100000000),('nan',None),('1.2万+',None),(True,None),('1e100',None)])
def test_number(value, expected):
    assert schema.number(value) == expected


@pytest.mark.parametrize('value', [1700000000,1700000000000,1700000000000000,'2023-11-15T06:13:20+08:00','2023-11-14T22:13:20Z'])
def test_timestamps(value):
    assert schema.timestamp(value) == '2023-11-14T22:13:20.000000Z'


def test_missing_and_naive_time():
    assert schema.timestamp('2026-01-01') == '2025-12-31T16:00:00.000000Z'
    assert schema.timestamp('20260101') == schema.timestamp('2026-01-01')
    for value in [None,0,'bad','99999999999999999999']:
        assert schema.timestamp(value) is None
    item = schema.normalize(note())
    assert item['favorites'] is None and item['comments'] is None
    with pytest.raises(ValueError): schema.normalize({'platform':'xhs'})


def test_read_does_not_create_or_migrate(client, database):
    assert client.get('/api/analytics/overview').json()['ready'] is False
    assert not database.exists()
    with sqlite3.connect(database) as db:
        db.execute('CREATE TABLE original (id INTEGER)')
    before = database.read_bytes()
    assert not client.get('/api/analytics/overview').json()['ready']
    assert database.read_bytes() == before


def test_migration_backup_duplicate_backfill_and_live_trigger(database):
    with sqlite3.connect(database) as db:
        db.execute('CREATE TABLE xhs_note (note_id TEXT,time INTEGER,last_modify_ts INTEGER,add_ts INTEGER,liked_count TEXT, nickname TEXT)')
        db.executemany('INSERT INTO xhs_note VALUES (?,?,?,?,?,?)', [('a',1700000000000,1700000010,1700000010,'10','匿名'), ('a',1700000000000,1700000020,1700000010,'20','匿名')])
    result = schema.migrate(database)
    assert Path(result['backup']).exists()
    with sqlite3.connect(result['backup']) as db:
        assert db.execute('SELECT count(*) FROM xhs_note').fetchone()[0] == 2
        assert not db.execute("SELECT name FROM sqlite_master WHERE name='analytics_content'").fetchone()
    with sqlite3.connect(database) as db:
        schema.register_functions(db)
        assert db.execute('SELECT likes FROM analytics_content').fetchone()[0] == 20
        assert db.execute('SELECT count(*) FROM analytics_snapshot').fetchone()[0] == 2
        db.execute("UPDATE xhs_note SET liked_count='30',last_modify_ts=1700000030 WHERE last_modify_ts=1700000020")
        assert db.execute('SELECT likes FROM analytics_content').fetchone()[0] == 30
    assert schema.migrate(database)['backup'] is None
    with sqlite3.connect(database) as db:
        assert db.execute('SELECT count(*) FROM analytics_snapshot').fetchone()[0] == 3
        assert db.execute('SELECT count(*) FROM xhs_note').fetchone()[0] == 2


@pytest.mark.asyncio
async def test_real_sqlalchemy_collector_insert_update_rollback(database, monkeypatch):
    monkeypatch.setattr(db_session, '_engines', {})
    import config
    monkeypatch.setattr(config, 'SAVE_DATA_OPTION', 'sqlite')
    await db_session.create_tables('sqlite')
    engine = db_session.get_async_engine('sqlite')
    try:
        async with engine.begin() as conn:
            await conn.execute(text("INSERT INTO xhs_note(note_id,time,last_modify_ts,liked_count) VALUES('a',1700000000,1700000010,'1万')"))
        async with engine.begin() as conn:
            await conn.execute(text("UPDATE xhs_note SET last_modify_ts=1700000020,liked_count='20000' WHERE note_id='a'"))
        with pytest.raises(RuntimeError):
            async with engine.begin() as conn:
                await conn.execute(text("UPDATE xhs_note SET last_modify_ts=1700000030,liked_count='30000' WHERE note_id='a'"))
                raise RuntimeError('rollback')
        with sqlite3.connect(database) as db:
            assert db.execute('SELECT likes FROM analytics_content').fetchone()[0] == 20000
            assert db.execute('SELECT count(*) FROM analytics_snapshot').fetchone()[0] == 2
    finally:
        await engine.dispose()


def test_snapshots_require_two_valid_observations(database):
    a = note()
    insert(database, [a])
    history = service.history('xhs','a')
    assert not history['metrics']['likes']['comparable']
    assert history['metrics']['likes']['rate'] is None
    insert(database, [dict(a, collected_at='2026-02-02', likes='0'),dict(a, collected_at='2026-02-03', likes='13000')])
    history = service.history('xhs','a')
    assert history['snapshot_count'] == 3
    assert history['metrics']['likes']['delta'] == 1000
    assert history['metrics']['favorites']['delta'] is None
    # New observations do not overwrite historical values; older imports do not
    # replace latest content even when imported later.
    insert(database, [dict(a, collected_at='2025-12-01', likes='0')])
    assert service.history('xhs','a')['metrics']['likes']['rate'] is None
    with service.read_db() as (db,_):
        assert db.execute('SELECT likes FROM analytics_content').fetchone()[0] == 13000


def test_filter_aggregate_rank_comment_grain_and_pagination(client,database):
    insert(database,[note(),dict(note('b'),platform='dy',author='第二作者',likes=None, keyword='美食'),dict(note('c'),author=None, published_at=None, collected_at=None, likes=None)])
    insert(database,[dict(platform='xhs',content_id='a',comment_id='c1',content='旅行很开心',create_time='2026-01-02',collected_at='2026-02-01'),dict(platform='xhs',content_id='a',comment_id='c2',content='旅行很开心',create_time='2026-01-02')],kind='comment')
    r=client.get('/api/analytics/overview').json()
    assert r['summary']['content_count']==3
    assert r['summary']['author_count']==2
    assert r['summary']['comment_count']==2
    assert r['summary']['valid_likes_count']==1
    assert r['summary']['valid_publish_count']==2
    assert r['summary']['valid_comments_count']==0
    assert sum(p['value'] for p in r['charts']['comment_trend'])==2
    assert r['charts']['words']
    assert client.get('/api/analytics/overview',params={'platform':'dy'}).json()['summary']['content_count']==1
    assert client.get('/api/analytics/overview',params={'keyword':'旅行'}).json()['summary']['content_count']==2
    assert client.get('/api/analytics/overview',params={'start':'2026-02-01','end':'2026-02-01','basis':'published_at'}).json()['summary']['content_count']==0
    assert client.get('/api/analytics/overview',params={'start':'2026-02-01','end':'2026-02-01','basis':'collected_at'}).json()['summary']['content_count']==2
    r=client.get('/api/analytics/details',params={'page_size':1,'sort':'interactions'}).json()
    assert r['total']==3 and r['items'][0]['content_id']=='a'
    assert r['items'][0]['valid_metric_count']==1 and r['items'][0]['interactions']==12000
    assert client.get('/api/analytics/details',params={'page':2,'page_size':1}).json()['page']==2
    assert client.get('/api/analytics/details',params={'kind':'comment','word':'开心','comment_date':'2026-01-02'}).json()['total']==2
    assert client.get('/api/analytics/details',params={'content_type':'__missing__'}).json()['total']==3
    for params in [{'start':'invalid'},{'basis':'SQL'},{'start':'2026-02-01','end':'2026-01-01'},{'platform':"';DROP TABLE x;--"}]:
        assert client.get('/api/analytics/overview',params=params).status_code==422
    assert client.get('/api/analytics/details?page_size=1000').status_code==422


@pytest.mark.parametrize('format',['csv','json','jsonl','xlsx'])
def test_import_formats_validation_dedup_and_export(client,database,format):
    client.post('/api/analytics/initialize')
    records=[note(),note(),{'platform':'xhs','title':'missing ID'},dict(note('b'),likes='unknown',published_at='broken')]
    if format=='json': data=json.dumps(records).encode()
    elif format=='jsonl': data='\n'.join(json.dumps(r) for r in records).encode()
    elif format=='csv':
        out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=list(note())+['title']);writer.writeheader();writer.writerows(records);data=out.getvalue().encode()
    else:
        wb=Workbook();ws=wb.active;keys=list(note())+['title'];ws.append(keys)
        for r in records: ws.append([r.get(k) for k in keys])
        out=io.BytesIO();wb.save(out);data=out.getvalue()
    params={'filename':'source.'+format}
    r=client.post('/api/analytics/import',params=params,content=data)
    assert r.status_code==200, r.text
    assert r.json()['inserted']==2 and r.json()['duplicates']==1 and r.json()['rejected']==1
    assert r.json()['warnings']==1
    assert client.get('/api/analytics/details').json()['total']==0
    r=client.post('/api/analytics/import',params={**params,'dry_run':'false'},content=data)
    assert r.status_code==200,r.text
    assert client.get('/api/analytics/details').json()['total']==2
    r=client.post('/api/analytics/import',params={**params,'dry_run':'false'},content=data).json()
    assert r['duplicates']==3
    exported=client.get('/api/analytics/export?platform=xhs&format=xlsx')
    wb=load_workbook(io.BytesIO(exported.content)); assert wb.active.max_row==3
    exported=client.get('/api/analytics/export?keyword=旅行&format=csv')
    assert len(list(csv.DictReader(io.StringIO(exported.content.decode('utf-8-sig')))))==2


def test_import_conflicting_same_time_and_missing_time(client,database):
    client.post('/api/analytics/initialize')
    data=[note(),dict(note(),likes='30000'),dict(note('b'),collected_at=None)]
    r=client.post('/api/analytics/import',params={'filename':'a.json','dry_run':'false'},json=data)
    assert r.status_code==200,r.text
    assert r.json()['rejected']==1
    assert service.history('xhs','a')['snapshot_count']==1
    assert service.history('xhs','b')['snapshot_count']==0
    before=database.read_bytes()
    assert client.get('/api/analytics/export').status_code==200
    assert client.get('/api/analytics/overview').status_code==200
    assert database.read_bytes()==before


def test_import_comment_aliases_and_bad_files(client,database):
    insert(database,[note()])
    r=client.post('/api/analytics/import',params={'filename':'a.jsonl','kind':'comment','platform':'xhs','dry_run':'false'},content=json.dumps({'note_id':'a','comment_id':'c','content':'原始评论','create_time':1700000000000}))
    assert r.status_code==200,r.text
    assert client.get('/api/analytics/details?kind=comment').json()['items'][0]['body']=='原始评论'
    for filename,data in [('a.json',b'bad'),('a.xls',b'bad'),('a.xlsx',b'bad')]:
        assert client.post('/api/analytics/import',params={'filename':filename},content=data).status_code==400


def test_csv_formula_and_null_export(client,database):
    insert(database,[dict(note(), title='=HYPERLINK("bad")')])
    r=client.get('/api/analytics/export')
    row=next(csv.DictReader(io.StringIO(r.content.decode('utf-8-sig'))))
    assert row['title'].startswith("'=")
    assert row['favorites']==''
