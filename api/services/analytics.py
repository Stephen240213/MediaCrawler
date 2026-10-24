"""Bounded imports and SQL-only aggregation over the collection SQLite file."""
import csv
import io
import json
import sqlite3
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path

from database.analytics import (SCHEMA, CONTENT_COLUMNS, PLATFORMS, ALIASES,
                                migrate, normalize, pick, save, timestamp)
from config.db_config import sqlite_db_config

MAX_ROWS = 50000
MAX_BYTES = 10 * 1024 * 1024


def db_path():
    return Path(sqlite_db_config['db_path'])


@contextmanager
def read_db():
    path = db_path()
    ready = False
    db = None
    try:
        if path.exists():
            db = sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True, timeout=10)
            ready = db.execute("SELECT 1 FROM sqlite_master WHERE name='analytics_meta'").fetchone() is not None
        if not ready:
            if db:
                db.close()
            db = sqlite3.connect(':memory:')
            for sql in SCHEMA:
                db.execute(sql)
        db.row_factory = sqlite3.Row
        db.execute('PRAGMA query_only=ON')
        yield db, ready
    finally:
        if db:
            db.close()


def filters(platform='', keyword='', start='', end='', basis='published_at', content_id='', author_id='', content_type=''):
    if platform and platform not in PLATFORMS:
        raise ValueError('无效平台')
    if basis not in ('published_at', 'collected_at'):
        raise ValueError('无效时间口径')
    clauses, args = ['1=1'], []
    for key, value in (('platform', platform), ('keyword', keyword), ('content_id', content_id), ('author_id', author_id), ('content_type', content_type)):
        if value:
            if value == '__missing__':
                clauses.append(f"nullif(c.{key},'') IS NULL")
            else:
                clauses.append(f'c.{key}=?')
                args.append(value)
    if start and end and start > end:
        raise ValueError('开始日期不能晚于结束日期')
    for value, op, next_day in ((start, '>=', False), (end, '<', True)):
        if value:
            try:
                dt = datetime.strptime(value, '%Y-%m-%d')
            except ValueError as exc:
                raise ValueError('日期须为 YYYY-MM-DD') from exc
            if next_day:
                dt += timedelta(days=1)
            clauses.append(f'c.{basis} {op} ?')
            args.append(timestamp(dt.isoformat()))
    return ' AND '.join(clauses), args


def rows(db, sql, args=()):
    return [dict(r) for r in db.execute(sql, args)]


def overview(options):
    where, args = filters(**options)
    with read_db() as (db, ready):
        # One read transaction ensures all panels describe the same snapshot.
        db.execute('BEGIN')
        summary = dict(db.execute(f'''SELECT count(*) content_count,
            count(DISTINCT CASE WHEN author_id IS NOT NULL THEN platform||':'||author_id END) author_count,
            max(collected_at) last_collected_at, count(published_at) valid_publish_count,
            count(collected_at) valid_collect_count, count(likes) valid_likes_count,
            count(favorites) valid_favorites_count, count(comments) valid_comments_count,
            count(shares) valid_shares_count FROM analytics_content c WHERE {where}''', args).fetchone())
        parent = f'''FROM analytics_comment m JOIN analytics_content c
            ON c.platform=m.platform AND c.content_id=m.content_id WHERE {where}'''
        summary['comment_count'] = db.execute('SELECT count(*) ' + parent, args).fetchone()[0]
        summary['valid_comment_time_count'] = db.execute('SELECT count(m.published_at) ' + parent, args).fetchone()[0]
        summary['data_updated_at'] = db.execute('SELECT max(t) FROM (SELECT max(collected_at) t FROM analytics_content UNION ALL SELECT max(collected_at) FROM analytics_comment)').fetchone()[0]
        charts = {}
        for key in ('published_at', 'collected_at'):
            # Display dates and date-range boundaries both use Asia/Shanghai.
            charts[key] = rows(db, f'''SELECT date(c.{key},'+8 hours') label,count(*) value FROM analytics_content c
                WHERE {where} AND c.{key} IS NOT NULL GROUP BY label ORDER BY label''', args)
        for key in ('platform', 'keyword', 'content_type'):
            charts[key] = rows(db, f'''SELECT coalesce(nullif(c.{key},''),'未提供') label,count(*) value
                FROM analytics_content c WHERE {where} GROUP BY label ORDER BY value DESC,label LIMIT 30''', args)
        charts['comment_trend'] = rows(db, 'SELECT date(m.published_at,\'+8 hours\') label,count(*) value ' + parent + ' AND m.published_at IS NOT NULL GROUP BY label ORDER BY label', args)
        # SQL groups repeated bodies first; bounded term extraction avoids loading
        # the whole comment database. UI discloses this sampling boundary.
        bodies = rows(db, 'SELECT m.body,count(*) n ' + parent + " AND m.body IS NOT NULL GROUP BY m.body ORDER BY max(coalesce(m.collected_at,m.published_at,'')) DESC,m.body LIMIT 5000", args)
        import jieba
        from collections import Counter
        import re
        stop = {'一个', '这个', '那个', '我们', '你们', '他们', '可以', '就是', '真的', 'the', 'and', 'that', 'with', 'this'}
        words = Counter()
        for row in bodies:
            # Per-document presence, weighted by distinct comment count.
            tokens = set(w.lower() for w in jieba.cut((row['body'] or '')[:2000]) if len(w.strip()) > 1 and re.search(r'[\u4e00-\u9fffA-Za-z]', w))
            for word in tokens - stop:
                words[word] += row['n']
        charts['words'] = [{'label': w, 'value': n} for w, n in words.most_common(30)]
        authors = rows(db, f'''SELECT c.platform,c.author_id,max(c.author) author,count(*) content_count,
            sum(c.likes) likes,count(c.likes) valid_likes_count FROM analytics_content c WHERE {where}
            AND c.author_id IS NOT NULL GROUP BY c.platform,c.author_id ORDER BY content_count DESC,c.platform,c.author_id LIMIT 20''', args)
        return dict(ready=ready, summary=summary, charts=charts, authors=authors,
                    generated_at=datetime.now(timezone.utc).isoformat(),
                    keywords=rows(db, "SELECT DISTINCT keyword FROM analytics_content WHERE keyword IS NOT NULL AND keyword!='' ORDER BY keyword LIMIT 500"),
                    word_sample_bodies=len(bodies))


INTERACTIONS = '(coalesce(c.likes,0)+coalesce(c.favorites,0)+coalesce(c.comments,0)+coalesce(c.shares,0))'


def details(options, page=1, page_size=20, kind='content', sort='recent', word='', comment_date=''):
    where, args = filters(**options)
    with read_db() as (db, _):
        if kind == 'comment':
            source = f'''FROM analytics_comment m JOIN analytics_content c ON c.platform=m.platform AND c.content_id=m.content_id WHERE {where}'''
            if word:
                source += " AND instr(lower(m.body),lower(?))>0"
                args.append(word)
            if comment_date:
                datetime.strptime(comment_date, '%Y-%m-%d')
                source += " AND date(m.published_at,'+8 hours')=?"
                args.append(comment_date)
            columns = 'm.*,c.title,c.url,c.keyword'
            order = 'm.published_at DESC,m.platform,m.comment_id'
        else:
            source = f'FROM analytics_content c WHERE {where}'
            columns = f'''c.*,CASE WHEN c.likes IS NULL AND c.favorites IS NULL AND c.comments IS NULL AND c.shares IS NULL
                THEN NULL ELSE {INTERACTIONS} END interactions,
                (c.likes IS NOT NULL)+(c.favorites IS NOT NULL)+(c.comments IS NOT NULL)+(c.shares IS NOT NULL) valid_metric_count'''
            order = ('interactions DESC,' if sort == 'interactions' else '') + 'c.collected_at DESC,c.platform,c.content_id'
        db.execute('BEGIN')
        total = db.execute('SELECT count(*) ' + source, args).fetchone()[0]
        items = rows(db, f'SELECT {columns} {source} ORDER BY {order} LIMIT ? OFFSET ?', [*args, page_size, (page-1)*page_size])
        return dict(items=items, total=total, page=page, page_size=page_size)


def history(platform, content_id):
    with read_db() as (db, _):
        valid = {}
        for metric in ('likes', 'favorites', 'comments', 'shares'):
            first = db.execute(f'''SELECT collected_at,{metric} value FROM analytics_snapshot
                WHERE platform=? AND content_id=? AND {metric} IS NOT NULL ORDER BY collected_at LIMIT 1''', (platform, content_id)).fetchone()
            last = db.execute(f'''SELECT collected_at,{metric} value FROM analytics_snapshot
                WHERE platform=? AND content_id=? AND {metric} IS NOT NULL ORDER BY collected_at DESC LIMIT 1''', (platform, content_id)).fetchone()
            comparable = first and last and first['collected_at'] != last['collected_at']
            valid[metric] = dict(comparable=bool(comparable), delta=last['value']-first['value'] if comparable else None,
                                 rate=(last['value']-first['value']) / first['value'] if comparable and first['value'] > 0 else None,
                                 first=dict(first) if first else None, last=dict(last) if last else None)
        return {'metrics': valid, 'snapshot_count': db.execute('SELECT count(*) FROM analytics_snapshot WHERE platform=? AND content_id=?', (platform, content_id)).fetchone()[0],
                'items': rows(db, 'SELECT * FROM analytics_snapshot WHERE platform=? AND content_id=? ORDER BY collected_at DESC LIMIT 100', (platform, content_id))}


def parse_file(name, data):
    suffix = Path(name).suffix.lower()
    if len(data) > MAX_BYTES:
        raise ValueError('文件不能超过 10 MB')
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        import zipfile
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Excel 解压后过大')
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if wb.active is None:
                raise ValueError('Excel 文件没有可读取的工作表')
            iterator = wb.active.iter_rows(values_only=True)
            header = next(iterator, None)
            if not header or any(h is None for h in header) or len(set(header)) != len(header):
                raise ValueError('Excel 首行需要非空且不重复的字段名')
            result = []
            for row in iterator:
                if len(result) >= MAX_ROWS:
                    raise ValueError('每次最多导入 50000 行')
                if any(v is not None for v in row):
                    result.append(dict(zip(header, row)))
            return result
        finally:
            wb.close()
    text = data.decode('utf-8-sig')
    if suffix == '.csv':
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames or any(not h.strip() for h in reader.fieldnames) or len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError('CSV 字段名不能为空或重复')
        result = []
        for row in reader:
            if len(result) >= MAX_ROWS:
                raise ValueError('每次最多导入 50000 行')
            result.append(row)
        return result
    if suffix == '.jsonl':
        result = [json.loads(line) for line in text.splitlines() if line.strip()]
    elif suffix == '.json':
        result = json.loads(text)
        if isinstance(result, dict):
            result = result.get('data', result.get('contents', result.get('comments', [result])))
    else:
        raise ValueError('支持 .xlsx、UTF-8 CSV、JSON、JSONL（旧 .xls 请另存为 .xlsx）')
    if not isinstance(result, list) or len(result) > MAX_ROWS:
        raise ValueError('文件需要记录数组，且每次最多 50000 行')
    return result


def import_file(name, data, platform='', kind='content', dry_run=True):
    source = parse_file(name, data)
    result = dict(total=len(source), inserted=0, updated=0, duplicates=0, rejected=0, warnings=0, errors=[], dry_run=dry_run)
    cleaned, seen, batch_current, batch_snapshots = [], set(), {}, {}
    with read_db() as (db, ready):
        if not ready:
            raise ValueError('请先初始化分析结构（会备份已有采集库）')
        for index, row in enumerate(source, 1):
            try:
                if not isinstance(row, dict):
                    raise ValueError('记录必须是对象')
                item = normalize(row, platform or None, kind)
                identity = item['content_id'] if kind == 'content' else item['comment_id']
                fingerprint = json.dumps(item, sort_keys=True, ensure_ascii=False)
                if fingerprint in seen:
                    result['duplicates'] += 1
                    continue
                seen.add(fingerprint)
                table = 'analytics_content' if kind == 'content' else 'analytics_comment'
                key = 'content_id' if kind == 'content' else 'comment_id'
                identity_key = (item['platform'], identity)
                previous = batch_current.get(identity_key)
                if previous is None:
                    previous = db.execute(f'SELECT * FROM {table} WHERE platform=? AND {key}=?', identity_key).fetchone()
                # Duplicates with a new observation time are meaningful snapshots.
                if previous and all(previous[k] == v for k, v in item.items()):
                    result['duplicates'] += 1
                    continue
                if kind == 'content' and item['collected_at']:
                    snapshot_key = (item['platform'], identity, item['collected_at'])
                    old = batch_snapshots.get(snapshot_key)
                    if old is None:
                        old = db.execute('SELECT * FROM analytics_snapshot WHERE platform=? AND content_id=? AND collected_at=?', snapshot_key).fetchone()
                    if old and any(old[k] != item[k] for k in ('likes', 'favorites', 'comments', 'shares')):
                        raise ValueError('同一内容与采集时间已有不同指标快照，请核对采集时间')
                    if old:
                        result['duplicates'] += 1
                        continue
                if previous and previous['collected_at'] and item['collected_at'] is None:
                    raise ValueError('已有记录含采集时间，当前记录缺少采集时间，不能安全覆盖')
                warnings = []
                for key in ('published_at', 'collected_at', 'likes', 'favorites', 'comments', 'shares'):
                    if kind == 'comment' and key not in ('published_at', 'collected_at'):
                        continue
                    raw = pick(row, ALIASES[key])
                    if raw is not None and item[key] is None:
                        warnings.append(key + ' 无法解析，保留为空')
                if warnings or item['collected_at'] is None:
                    result['warnings'] += 1
                    if len(result['errors']) < 20:
                        result['errors'].append({'row': index, 'warning': '；'.join(warnings or ['缺少采集时间，不生成快照'])})
                result['updated' if previous else 'inserted'] += 1
                cleaned.append(item)
                if previous is None or (item['collected_at'] or '') >= (previous['collected_at'] or ''):
                    batch_current[identity_key] = item
                if kind == 'content' and item['collected_at']:
                    batch_snapshots[(item['platform'], identity, item['collected_at'])] = item
            except ValueError as exc:
                result['rejected'] += 1
                if len(result['errors']) < 20:
                    result['errors'].append({'row': index, 'error': str(exc)})
    if not dry_run:
        # Serialize writes against other imports/collector writes.
        # Recheck collisions after acquiring lock to preserve immutable history.
        with sqlite3.connect(db_path(), timeout=30) as db:
            db.execute('BEGIN IMMEDIATE')
            for item in cleaned:
                if kind == 'content' and item['collected_at']:
                    old = db.execute('SELECT likes,favorites,comments,shares FROM analytics_snapshot WHERE platform=? AND content_id=? AND collected_at=?', (item['platform'], item['content_id'], item['collected_at'])).fetchone()
                    if old and tuple(item[k] for k in ('likes', 'favorites', 'comments', 'shares')) != old:
                        raise ValueError('导入批次含冲突快照，未写入任何记录；请修正后重试')
                save(db, item, kind)
    return result
